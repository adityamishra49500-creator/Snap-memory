"""
Document ingestion pipeline: FILE -> validation -> text extraction ->
cleaning -> chunking -> metadata -> (caller handles embeddings + storage).
"""
import os
import re
from core import config

SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".ppt", ".txt", ".md"}


class UnsupportedFileError(Exception):
    pass


class ExtractionError(Exception):
    pass


def validate_file(filename: str, data: bytes, max_bytes: int = 200 * 1024 * 1024):
    ext = os.path.splitext(filename)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileError(f"Unsupported file type: {ext or '(no extension)'}")
    if not data:
        raise ExtractionError("File is empty.")
    if len(data) > max_bytes:
        raise ExtractionError(f"File exceeds max size of {max_bytes} bytes.")
    # basic path traversal protection for the filename we persist to disk
    safe_name = os.path.basename(filename)
    if safe_name != filename:
        filename = safe_name
    return ext, filename


def clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf(path: str):
    """Returns list of (page_number, text). Tries pdfplumber first (better
    layout handling), falls back to pypdf if pdfplumber fails on a
    malformed file."""
    pages = []
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                try:
                    text = page.extract_text() or ""
                except Exception:
                    text = ""
                pages.append((i, clean_text(text)))
        if any(t for _, t in pages):
            return pages
    except Exception:
        pass

    try:
        from pypdf import PdfReader
        reader = PdfReader(path)
        pages = []
        for i, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            pages.append((i, clean_text(text)))
        return pages
    except Exception as e:
        raise ExtractionError(f"Could not extract text from PDF (corrupted or encrypted?): {e}")


def extract_pptx(path: str):
    """Returns list of (slide_number, text)."""
    try:
        from pptx import Presentation
    except ImportError:
        raise ExtractionError("python-pptx is not installed.")
    try:
        prs = Presentation(path)
    except Exception as e:
        raise ExtractionError(f"Could not open PPTX (corrupted?): {e}")

    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        texts.append(line)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text for cell in row.cells)
                    if row_text.strip():
                        texts.append(row_text)
        slides.append((i, clean_text("\n".join(texts))))
    return slides


def extract_txt_or_md(path: str):
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            return clean_text(f.read())
    except Exception as e:
        raise ExtractionError(f"Could not read text file: {e}")


def chunk_text(text: str, chunk_size: int = None, overlap: int = None):
    """Simple, deterministic sliding-window chunker on paragraph/sentence
    boundaries where possible. Returns a list of chunk strings."""
    chunk_size = chunk_size or config.CHUNK_SIZE_CHARS
    overlap = overlap or config.CHUNK_OVERLAP_CHARS
    text = text.strip()
    if not text:
        return []
    if len(text) <= chunk_size:
        return [text]

    # Split into sentence-ish units first so we don't cut mid-sentence when avoidable.
    units = re.split(r"(?<=[.!?])\s+|\n+", text)
    units = [u for u in units if u.strip()]

    chunks = []
    current = ""
    for unit in units:
        if len(current) + len(unit) + 1 <= chunk_size:
            current = f"{current} {unit}".strip()
        else:
            if current:
                chunks.append(current)
            # start new chunk, carrying overlap from the tail of the previous one
            tail = current[-overlap:] if overlap and current else ""
            current = f"{tail} {unit}".strip()
    if current:
        chunks.append(current)

    # Guard against any single oversized unit (e.g. a huge unformatted blob)
    final = []
    for c in chunks:
        if len(c) <= chunk_size * 1.5:
            final.append(c)
        else:
            for i in range(0, len(c), chunk_size - overlap):
                final.append(c[i:i + chunk_size])
    return final


def extract_and_chunk(ext: str, path: str):
    """Top-level entry point. Returns a list of dicts:
    {text, page_number, slide_number, chunk_index}
    """
    results = []
    if ext == ".pdf":
        pages = extract_pdf(path)
        idx = 0
        for page_num, text in pages:
            for chunk in chunk_text(text):
                results.append({"text": chunk, "page_number": page_num,
                                 "slide_number": None, "chunk_index": idx})
                idx += 1
    elif ext in (".pptx", ".ppt"):
        slides = extract_pptx(path)
        idx = 0
        for slide_num, text in slides:
            for chunk in chunk_text(text):
                results.append({"text": chunk, "page_number": None,
                                 "slide_number": slide_num, "chunk_index": idx})
                idx += 1
    elif ext in (".txt", ".md"):
        text = extract_txt_or_md(path)
        idx = 0
        for chunk in chunk_text(text):
            results.append({"text": chunk, "page_number": None,
                             "slide_number": None, "chunk_index": idx})
            idx += 1
    else:
        raise UnsupportedFileError(ext)

    if not results:
        raise ExtractionError("No extractable text found in this file.")
    return results
